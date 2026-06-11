"""Embedded metadata dates: PDF/DOCX/XLSX carry their own created/modified
dates that survive email downloads (unlike filesystem mtime)."""
from datetime import datetime, timezone

import docx as python_docx
import openpyxl
import pymupdf

from processors import filedates

MODIFIED = datetime(2026, 1, 15, 9, 30, 0, tzinfo=timezone.utc)
CREATED = datetime(2025, 11, 3, 14, 0, 0, tzinfo=timezone.utc)


def test_docx_prefers_modified_over_created(tmp_path):
    path = tmp_path / "doc.docx"
    document = python_docx.Document()
    document.add_paragraph("hello")
    document.core_properties.created = CREATED
    document.core_properties.modified = MODIFIED
    document.save(str(path))

    assert filedates.embedded_date("docx", str(path)) == MODIFIED


def test_xlsx_reads_workbook_properties(tmp_path):
    # openpyxl stamps `modified` at save time (as Excel does), so compare
    # against what actually landed in the file
    path = tmp_path / "sheet.xlsx"
    openpyxl.Workbook().save(str(path))
    stored = openpyxl.load_workbook(str(path)).properties.modified

    assert filedates.embedded_date("spreadsheet", str(path)) == stored.replace(tzinfo=timezone.utc)


def test_pdf_parses_metadata_date_with_offset(tmp_path):
    path = tmp_path / "doc.pdf"
    document = pymupdf.open()
    document.new_page()
    document.set_metadata({"modDate": "D:20260115133000+04'00'"})
    document.save(str(path))
    document.close()

    # 13:30 at UTC+4 is 09:30 UTC
    assert filedates.embedded_date("pdf", str(path)) == MODIFIED


def test_pdf_date_string_variants():
    assert filedates._parse_pdf_date("D:20260115093000Z") == MODIFIED
    assert filedates._parse_pdf_date("D:20260115") == datetime(2026, 1, 15, tzinfo=timezone.utc)
    assert filedates._parse_pdf_date("") is None
    assert filedates._parse_pdf_date(None) is None
    assert filedates._parse_pdf_date("garbage") is None


def test_unreadable_file_returns_none(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a real docx")

    assert filedates.embedded_date("docx", str(path)) is None


def test_plaintext_has_no_embedded_date(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("hello")

    assert filedates.embedded_date("plaintext", str(path)) is None
